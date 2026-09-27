"""Excel çıktısı — üç sayfa: Çizelge · Özet · Eksikler."""

from datetime import date, timedelta
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from app.repositories import drafts as drafts_repo
from app.repositories import schedule as repo
from app.routers.schedule import GUN_ADI, AY_ADI, VARDIYA_ETIKET

router = APIRouter(tags=["çıktı"])

# DESIGN §2 renkleri — ekranla aynı dil
BRAND = "1D265F"
BEYAZ = "FFFFFF"
ZEMIN = "F5F7FA"
CIZGI = "E3E8EF"
KIRMIZI = "C62828"

GOREV_KISA = {"TRIYAJ": "TRY", "AMBULANS": "AMB", "GOZLEM": "GÖZ"}
IZIN_ADI = {"yillik_izin": "İzin", "rapor": "Rapor", "ucretsiz_izin": "İzin", "diger": "İzin"}

ince = Side(style="thin", color=CIZGI)
KENARLIK = Border(left=ince, right=ince, top=ince, bottom=ince)


def _buyuk(metin: str) -> str:
    """Türkçe büyük harf. Python'un upper()'ı 'i' → 'I' yapıyor, 'İ' olmalı:
    "Ekip Liderleri" → "EKIP LIDERLERI" değil "EKİP LİDERLERİ"."""
    return metin.replace("i", "İ").replace("ı", "I").upper()


def _tarih(d: date) -> str:
    return f"{d.day} {AY_ADI[d.month - 1]} {d.year}"


def _dosya_adi(ad: str, bas: date, bitis: date) -> str:
    temiz = "".join(c if c.isalnum() or c in " -_" else "" for c in ad).strip().replace(" ", "_")
    return f"Acibadem_Nobet_{temiz}_{bas:%Y%m%d}-{(bitis - timedelta(days=1)):%Y%m%d}.xlsx"


@router.get("/drafts/{draft_id}/export.xlsx", summary="Çizelgeyi Excel'e aktar")
async def excel(draft_id: int) -> StreamingResponse:
    if (t := await repo.taslak(draft_id)) is None:
        raise HTTPException(status_code=404, detail="Taslak bulunamadı.")

    bas: date = t["period_start"]
    bitis: date = t["period_end"]          # DIŞLAYICI
    son = bitis - timedelta(days=1)
    veri = await repo.cizelge_verisi(draft_id, bas, son)

    gunler = [bas + timedelta(days=i) for i in range((bitis - bas).days)]

    kitap = Workbook()
    _sayfa_cizelge(kitap.active, t, gunler, veri)
    _sayfa_ozet(kitap.create_sheet("Özet"), veri, bas, bitis)
    _sayfa_eksikler(kitap.create_sheet("Eksikler"), veri)

    akis = BytesIO()
    kitap.save(akis)
    akis.seek(0)

    ad = _dosya_adi(t["name"], bas, bitis)
    return StreamingResponse(
        akis,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        # Türkçe karakterler için RFC 5987: filename* alanı UTF-8 taşır
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(ad)}"},
    )


def _sayfa_cizelge(ws, t: dict, gunler: list[date], veri: dict) -> None:
    ws.title = "Çizelge"
    ws["A1"] = t["name"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{_tarih(gunler[0])} – {_tarih(gunler[-1])} · {t['unit_name']}"
    ws["A2"].font = Font(color="5B6B7F")

    BASLIK = 4
    ws.cell(BASLIK, 1, "Personel").font = Font(bold=True)
    for i, g in enumerate(gunler):
        h = ws.cell(BASLIK, 2 + i, f"{GUN_ADI[g.weekday()]}\n{g.day} {AY_ADI[g.month - 1]}")
        h.font = Font(bold=True, size=9)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.fill = PatternFill("solid", fgColor=ZEMIN)
        h.border = KENARLIK

    atamalar = {(a["staff_id"], a["work_date"]): a for a in veri["atamalar"]}
    izinler: dict[tuple[int, date], str] = {}
    for iz in veri["izinler"]:
        g = iz["period"].lower
        while g < iz["period"].upper:
            izinler[(iz["staff_id"], g)] = iz["absence_type"]
            g += timedelta(days=1)

    # Rol gruplarıyla — ekrandaki düzenin aynısı
    gruplar = [
        ("Sorumlu & Eğitim", {"sorumlu_hemsire", "egitim_hemsire"}),
        ("Ekip Liderleri", {"shift_yetkilisi"}),
        ("Hemşireler", None),
    ]
    satir = BASLIK + 1
    for ad, roller in gruplar:
        ozel = {"sorumlu_hemsire", "egitim_hemsire", "shift_yetkilisi"}
        kisiler = [
            k for k in veri["personel"]
            if (k["role_code"] in roller) if roller is not None
        ] if roller is not None else [
            k for k in veri["personel"] if k["role_code"] not in ozel
        ]
        if not kisiler:
            continue
        b = ws.cell(satir, 1, _buyuk(ad))
        b.font = Font(bold=True, size=9, color="5B6B7F")
        b.fill = PatternFill("solid", fgColor=ZEMIN)
        for c in range(2, len(gunler) + 2):
            ws.cell(satir, c).fill = PatternFill("solid", fgColor=ZEMIN)
        satir += 1

        for k in kisiler:
            ws.cell(satir, 1, k["full_name"]).border = KENARLIK
            for i, g in enumerate(gunler):
                h = ws.cell(satir, 2 + i)
                h.border = KENARLIK
                h.alignment = Alignment(horizontal="center", vertical="center")
                a = atamalar.get((k["id"], g))
                if a is not None:
                    etiket = VARDIYA_ETIKET.get(a["shift_code"], "G")
                    gorevler = [GOREV_KISA.get(x, x) for x in (a["tasks"] or [])]
                    h.value = f"{etiket} · {' '.join(gorevler)}" if gorevler else etiket
                    if etiket == "N":
                        h.fill = PatternFill("solid", fgColor=BRAND)
                        h.font = Font(color=BEYAZ, bold=True, size=9)
                    else:
                        h.font = Font(color=BRAND, bold=True, size=9)
                elif (k["id"], g) in izinler:
                    h.value = IZIN_ADI.get(izinler[(k["id"], g)], "İzin")
                    h.font = Font(color="5B6B7F", size=9)
            satir += 1

    ws.column_dimensions["A"].width = 24
    for i in range(len(gunler)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 11
    ws.row_dimensions[BASLIK].height = 30
    ws.freeze_panes = ws.cell(BASLIK + 1, 2)       # ilk sütun + başlık dondurulur

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.print_title_rows = f"{BASLIK}:{BASLIK}"


def _sayfa_ozet(ws, veri: dict, bas: date, bitis: date) -> None:
    basliklar = ["Personel", "Rol", "Toplam saat", "Gece", "Hafta sonu", "Hedef", "Fark"]
    for i, b in enumerate(basliklar, start=1):
        h = ws.cell(1, i, b)
        h.font = Font(bold=True)
        h.fill = PatternFill("solid", fgColor=ZEMIN)
        h.border = KENARLIK

    gun_sayisi = (bitis - bas).days
    ay_gun = (date(bas.year + (bas.month == 12), bas.month % 12 + 1, 1) - date(bas.year, bas.month, 1)).days

    aylik = {m["staff_id"]: m for m in veri["aylik"]}
    gece = {}
    haftasonu = {}
    for a in veri["atamalar"]:
        if a["crosses_midnight"]:
            gece[a["staff_id"]] = gece.get(a["staff_id"], 0) + 1
        if a["work_date"].weekday() >= 5:
            haftasonu[a["staff_id"]] = haftasonu.get(a["staff_id"], 0) + 1

    satir = 2
    for k in veri["personel"]:
        m = aylik.get(k["id"])
        saat = float(m["planned_hours"]) if m else 0.0
        hedef = round(float(m["min_hours"]) * gun_sayisi / ay_gun, 1) if m and ay_gun else 0.0
        fark = round(saat - hedef, 1)
        for i, deger in enumerate(
            [k["full_name"], k["role_name"], saat, gece.get(k["id"], 0),
             haftasonu.get(k["id"], 0), hedef, fark], start=1,
        ):
            h = ws.cell(satir, i, deger)
            h.border = KENARLIK
            if i == 7 and fark < 0:
                h.font = Font(color=KIRMIZI)
        satir += 1

    for kolon, genislik in zip("ABCDEFG", [24, 24, 12, 8, 12, 10, 10]):
        ws.column_dimensions[kolon].width = genislik
    ws.freeze_panes = "A2"


def _sayfa_eksikler(ws, veri: dict) -> None:
    basliklar = ["Gün", "Vardiya", "Görev", "Atanan", "Gereken", "Eksik"]
    for i, b in enumerate(basliklar, start=1):
        h = ws.cell(1, i, b)
        h.font = Font(bold=True)
        h.fill = PatternFill("solid", fgColor=ZEMIN)
        h.border = KENARLIK

    VARDIYA_ADI = {"GUNDUZ": "Gündüz", "GECE": "Gece", "GUNDUZ_CMT": "Gündüz (Cmt)"}
    SLOT_ADI = {"GENEL": "Genel mevcut", "TRIYAJ": "Triyaj", "GOZLEM": "Gözlem",
                "AMBULANS": "Ambulans", "SAYIM": "Sayım yetkilisi",
                "SHIFT_YETKILISI": "Ekip lideri"}

    satir = 2
    for k in sorted(veri["kapsama"], key=lambda x: (x["day"], x["shift_code"], x["slot_code"])):
        if k["assigned"] >= k["required"]:
            continue
        degerler = [
            _tarih(k["day"]), VARDIYA_ADI.get(k["shift_code"], k["shift_code"]),
            SLOT_ADI.get(k["slot_code"], k["slot_code"]),
            k["assigned"], k["required"], k["required"] - k["assigned"],
        ]
        for i, d in enumerate(degerler, start=1):
            h = ws.cell(satir, i, d)
            h.border = KENARLIK
            if i == 6:
                h.font = Font(color=KIRMIZI, bold=True)
        satir += 1

    if satir == 2:
        ws.cell(2, 1, "Eksik yok.")

    for kolon, genislik in zip("ABCDEF", [16, 14, 18, 10, 10, 10]):
        ws.column_dimensions[kolon].width = genislik
    ws.freeze_panes = "A2"
